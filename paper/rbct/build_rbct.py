"""Build the Boletín de Ciencias de la Tierra submission (Word) from paper/main.tex.

Usage (from the repository root, with python-docx and latex2mathml installed):

    uv pip install python-docx latex2mathml
    uv run python paper/rbct/build_rbct.py

Outputs in paper/rbct/: manuscript_blind.docx (no author data, for review),
manuscript_full.docx (with author data) and cover_letter.docx. Equations are written as native
Word equations when Microsoft Office's MML2OMML.XSL is available.
"""

from __future__ import annotations

import copy
import os
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT / "paper"
OUT = PAPER / "rbct"
TEMPLATE = OUT / "templates" / "rbct_article_template.docx"
LETTER_TEMPLATE = OUT / "templates" / "rbct_cover_letter_template.docx"
XSL = Path("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL")

AUTHOR = "Albert Johan Herrera"
AFFILIATION = "Programa de Ingeniería Topográfica y Geomática, Universidad del Quindío, Armenia, Colombia"
EMAIL = "alberthj.herreras@uqvirtual.edu.co"
ORCID = "0009-0009-1490-9075"
PHONE = os.environ.get("RBCT_AUTHOR_PHONE", "____________")  # kept out of the public repository
REPO = "https://github.com/alberttkd/static-gravity-anomalies-colombia"
BLIND_REPO = "[repository URL withheld for blind review]"

TITLE_EN = (
    "A reproducible, test-driven workflow for reducing and gridding legacy gravity compilations with "
    "calibrated uncertainty, verified on a synthetic benchmark"
)
TITLE_ES = (
    "Flujo reproducible y verificado por pruebas para reducir e interpolar compilaciones gravimétricas "
    "heredadas con incertidumbre calibrada, validado con una campaña sintética"
)

ABSTRACT_EN = (
    "Bouguer anomaly maps are often built from legacy gravity compilations whose reduction choices and gridding "
    "methods are undocumented and lack uncertainty. We present an open, test-driven workflow that produces "
    "free-air, simple and complete Bouguer anomalies of known quality, and verify it on a synthetic survey with "
    "known truth. A term-by-term reconciliation recovers the defects injected into two simulated legacy "
    "spreadsheets: a 2300 kg m⁻³ reduction density (up to 4.1 mGal) and a sign-inverted curvature term (bias up "
    "to 0.75 mGal). Nettleton and Parasnis estimators recover a known density in high relief but fail, and are "
    "flagged, when the anomaly follows a low-relief topography. Under nested spatial block cross-validation, "
    "kriging, equivalent sources and a spline reduce the error of inverse-distance weighting by 62–72%, whereas "
    "random cross-validation understates errors by factors of 1.7–5.1. Only kriging provides uncertainty, which "
    "reaches nominal coverage on held-out blocks (0.956) and against the truth (0.953)."
)
ABSTRACT_ES = (
    "Los mapas de anomalía de Bouguer suelen construirse con compilaciones heredadas cuyas decisiones de reducción "
    "e interpolación no están documentadas. Se presenta un flujo abierto que produce anomalías de aire libre y de "
    "Bouguer de calidad conocida, validado con una campaña sintética de verdad conocida. La conciliación término a "
    "término recupera los defectos de dos hojas de cálculo heredadas simuladas: una densidad de 2300 kg m⁻³ (hasta "
    "4,1 mGal) y una curvatura con signo invertido (sesgo de hasta 0,75 mGal). Los estimadores de Nettleton y "
    "Parasnis recuperan la densidad con alto relieve, pero fallan, y se marcan, cuando la anomalía sigue una "
    "topografía de bajo relieve. Con validación cruzada espacial anidada, el kriging, las fuentes equivalentes y un "
    "spline reducen el error de IDW entre 62 y 72 %; la validación aleatoria lo subestima entre 1,7 y 5,1 veces. "
    "Solo el kriging entrega incertidumbre con cobertura nominal (0,956 y 0,953)."
)
KEYWORDS_EN = [
    "Bouguer anomaly",
    "gravity reduction",
    "spatial cross-validation",
    "kriging",
    "equivalent sources",
    "uncertainty",
    "synthetic benchmark",
    "reproducibility",
]
KEYWORDS_ES = [
    "anomalía de Bouguer",
    "reducción gravimétrica",
    "validación cruzada espacial",
    "kriging",
    "fuentes equivalentes",
    "incertidumbre",
    "campaña sintética",
    "reproducibilidad",
]

AI_STATEMENT = (
    "During the preparation of this work the author used Claude (Anthropic), an AI assistant, to support software "
    "development and testing and the drafting and editing of the text. The author reviewed and edited all content "
    "and takes full responsibility for the publication."
)
BIO = (
    "A.J. Herrera is affiliated with the Topographic and Geomatics Engineering programme of the Universidad del "
    "Quindío, Armenia, Colombia. Research interests include physical geodesy, the processing of gravity data, "
    "geostatistics and reproducible open-source geospatial software."
)

# ---------------------------------------------------------------------------------------------------------------
# Bibliography
# ---------------------------------------------------------------------------------------------------------------

ACCENTS = {"'": "\u0301", "`": "\u0300", '"': "\u0308", "~": "\u0303", "^": "\u0302", "H": "\u030b", "c": "\u0327"}


def detex(s: str) -> str:
    import unicodedata

    s = s.replace("\\ss", "ß").replace("\\&", "&").replace("--", "–").replace("\\i", "i")
    s = re.sub(r"\{\\([`'\"~^Hc])\s*\{?(\w)\}?\}", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = re.sub(r"\\([`'\"~^])\{?(\w)\}?", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("`", "‘").replace("'", "’")
    return unicodedata.normalize("NFC", re.sub(r"\s+", " ", s)).strip()


def parse_bib(text: str) -> dict[str, dict[str, str]]:
    entries: dict[str, dict[str, str]] = {}
    for m in re.finditer(r"@(\w+)\{(\w+),(.*?)\n\s*\n", text + "\n\n", flags=re.S):
        kind, key, body = m.group(1).lower(), m.group(2), m.group(3)
        fields: dict[str, str] = {"type": kind}
        i = 0
        while True:
            fm = re.compile(r"(\w+)\s*=\s*\{").search(body, i)
            if not fm:
                break
            depth, j = 1, fm.end()
            while depth:
                depth += {"{": 1, "}": -1}.get(body[j], 0)
                j += 1
            fields[fm.group(1).lower()] = body[fm.end() : j - 1]
            i = j
        entries[key] = fields
    return entries


def initials(given: str) -> str:
    parts = []
    for tok in given.replace(".", ". ").split():
        sub = [t for t in tok.split("-") if t]
        parts.append("-".join(t[0] + "." for t in sub))
    return "".join(parts)


def split_authors(raw: str) -> list[tuple[str, str]]:
    out = []
    for a in re.split(r"\s+and\s+", raw):
        a = a.strip()
        if a.startswith("{{") or (a.startswith("{") and a.endswith("}") and "," not in a):
            out.append((detex(a), ""))
        elif "," in a:
            fam, given = a.split(",", 1)
            out.append((detex(fam), detex(given)))
        else:
            *given, fam = a.split()
            out.append((detex(fam), detex(" ".join(given))))
    return out


def cite_name(e: dict[str, str]) -> str:
    authors = split_authors(e["author"])
    if len(authors) == 1:
        return authors[0][0]
    if len(authors) == 2:
        return f"{authors[0][0]} and {authors[1][0]}"
    return f"{authors[0][0]} et al."


def format_reference(e: dict[str, str]) -> str:
    authors = split_authors(e["author"])
    names = [f"{fam}, {initials(given)}" if given else fam for fam, given in authors]
    if len(names) > 10:
        names = names[:6] + ["et al."]
        who = ", ".join(names)
    elif len(names) == 1:
        who = names[0]
    else:
        who = ", ".join(names[:-1]) + ", and " + names[-1]
    title = detex(e["title"])
    year = e.get("year", "")
    doi = f" DOI: https://doi.org/{e['doi']}" if "doi" in e else ""
    if e["type"] == "article":
        vol = e.get("volume", "")
        num = f"({detex(e['number'])})" if "number" in e else ""
        pages = detex(e.get("pages", ""))
        pp = (f", pp. {pages.replace('–', '-')}" if "–" in pages else f", {pages}") if pages else ""
        return f"{who}, {title}. {detex(e['journal'])}, {vol}{num}{pp}, {year}.{doi}"
    if e["type"] == "book":
        ed = f", {e['edition']}nd ed." if e.get("edition") == "2" else (f", {e['edition']}th ed." if "edition" in e else "")
        return f"{who}, {title}{ed}, {detex(e['publisher'])}, {year}.{doi}"
    extra = ", ".join(detex(e[k]) for k in ("publisher", "howpublished") if k in e)
    note = f" {detex(e['note'])}." if "note" in e else ""
    return f"{who}, {title}, {extra}, {year}.{note}{doi}"


# ---------------------------------------------------------------------------------------------------------------
# LaTeX body -> blocks
# ---------------------------------------------------------------------------------------------------------------

UNITS = {"": "", "m": "m", "km": "km", "mGal": "mGal", "km^2": "km²", "mGal^2": "mGal²"}
SUP = str.maketrans("0123456789-+", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺")


class Citations:
    def __init__(self, bib: dict[str, dict[str, str]]):
        self.bib, self.order = bib, []  # type: ignore[var-annotated]

    def num(self, key: str) -> int:
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def bracket(self, keys: list[str]) -> str:
        nums = sorted(self.num(k) for k in keys)
        groups, start = [], nums[0]
        for a, b in zip(nums, [*nums[1:], None], strict=True):
            if b != a + 1:
                groups.append(str(start) if start == a else (f"{start},{a}" if a == start + 1 else f"{start}-{a}"))
                start = b
        return "[" + ",".join(groups) + "]"


def simple_math(m: str) -> str | None:
    """Unicode rendering of trivial inline math, or None if it needs a real equation."""
    s = m.replace("\\mgal", " mGal").replace("\\kgm", " kg m⁻³").replace("\\,", "\u2009").replace("\\times", "×")
    s = s.replace("^\\circ", "°").replace("\\pm", "±").replace("\\geq", "≥").replace("<", "<").replace("\\%", "%")
    s = re.sub(r"\^\{(-?\d+)\}", lambda k: k.group(1).translate(SUP), s)
    s = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", s)
    s = re.sub(r"(?<![A-Za-z])-(?=\d)", "−", s)
    if re.search(r"[\\^_{}]", s):
        return None
    return s


def latex_text(s: str, cites: Citations, refs: dict[str, str]) -> list[tuple[str, str]]:
    """Convert a LaTeX paragraph into a list of ('t', text) and ('m', latex math) segments."""
    s = re.sub(r"\s+", " ", s.strip())
    s = re.sub(r"(Fig\.|Table)~\\ref\{([^}]+)\}", lambda m: f"{m.group(1)} {refs[m.group(2)]}", s)
    s = s.replace("Fig. ", "Fig.\u00a0").replace("Table ", "Table\u00a0")
    s = re.sub(
        r"\\citet\{([^}]+)\}",
        lambda m: " ".join([cite_name(cites.bib[k]) + " " + cites.bracket([k]) for k in m.group(1).split(",")]),
        s,
    )
    s = re.sub(r"\\cite(?:p|alp)\{([^}]+)\}", lambda m: cites.bracket(m.group(1).split(",")), s)
    s = re.sub(r"\(\[(\d[\d,\-]*)\]\)", r"[\1]", s)
    s = re.sub(r"\$[^$]+\$", lambda m: re.sub(r"\\SI\{([^}]*)\}\{([^}]*)\}", r"\1\\,\\mathrm{\2}", m.group(0)), s)
    s = re.sub(
        r"\\SI\{([^}]*)\}\{([^}]*)\}",
        lambda m: m.group(1).replace("\\,", "\u2009") + ("\u00a0" + UNITS[m.group(2)] if m.group(2) else ""),
        s,
    )
    s = re.sub(r"\\(?:texttt|emph)\{([^}]*)\}", r"\1", s)
    s = s.replace("``", "“").replace("''", "”").replace("---", "—").replace("--", "–")
    s = s.replace("\\%", "%").replace("e.g.\\ ", "e.g. ").replace("i.e.\\ ", "i.e. ")
    s = re.sub(r"\\'(\w)", lambda m: detex("\\'" + m.group(1)), s)
    parts = re.split(r"(\$[^$]+\$)", s)
    out: list[tuple[str, str]] = []
    for p in parts:
        if p.startswith("$"):
            m = p[1:-1]
            simple = simple_math(m)
            if simple is not None:
                out.append(("t", simple))
            else:
                out.append(("m", m.replace("\\mgal", "\\,\\mathrm{mGal}").replace("\\kgm", "\\,\\mathrm{kg\\,m^{-3}}")))
        else:
            p = (
                p.replace("\\mgal", "\u00a0mGal")
                .replace("\\kgm{}", "\u00a0kg\u00a0m⁻³")
                .replace("\\kgm", "\u00a0kg\u00a0m⁻³")
            )
            p = p.replace("\\,", "\u2009").replace("~", "\u00a0").replace("{}", "")
            out.append(("t", p))
    for kind, val in out:
        if kind == "t" and "\\" in val:
            raise ValueError(f"unconverted LaTeX: {val[:120]}")
    return out


def parse_body(tex: str) -> list[dict]:
    body = tex[tex.index("\\section{Introduction}") : tex.index("\\section*{Data and code availability}")]
    blocks: list[dict] = []
    pos = 0
    env_re = re.compile(r"\\begin\{(figure|table|equation|align)\}(.*?)\\end\{\1\}", re.S)
    for m in env_re.finditer(body):
        blocks += _text_blocks(body[pos : m.start()])
        kind, inner = m.group(1), m.group(2)
        if kind == "figure":
            w = re.search(r"\\includegraphics\[width=([\d.]*)\\textwidth\]\{([^}]+)\}", inner)
            blocks.append(
                {
                    "kind": "figure",
                    "file": w.group(2).replace(".pdf", ".png"),
                    "wide": (w.group(1) or "1") == "1" or float(w.group(1) or 1) > 0.8,
                    "caption": _braced(inner, "\\caption"),
                    "label": _braced(inner, "\\label"),
                }
            )
        elif kind == "table":
            rows = []
            tab = inner[inner.index("\\begin{tabular}") : inner.index("\\end{tabular}")]
            tab = tab.split("}", 2)[2] if tab.count("}") >= 2 else tab
            for line in tab.split("\\\\"):
                line = line.strip()
                rule = None
                for r in ("\\toprule", "\\midrule", "\\bottomrule"):
                    if line.startswith(r):
                        rule, line = r, line[len(r) :].strip()
                if line:
                    line = re.sub(
                        r"\\multicolumn\{(\d)\}\{[^}]*\}\{([^}]*)\}",
                        lambda m: m.group(2) + " &" * (int(m.group(1)) - 1),
                        line,
                    )
                    rows.append({"cells": [c.strip() for c in line.split("&")], "rule_before": rule})
                elif rule == "\\midrule" and rows:
                    rows[-1]["rule_after"] = True
            blocks.append(
                {"kind": "table", "rows": rows, "caption": _braced(inner, "\\caption"), "label": _braced(inner, "\\label")}
            )
        else:
            for eq in inner.strip().split("\\\\"):
                blocks.append({"kind": "equation", "latex": eq.replace("&", "").strip()})
        pos = m.end()
    blocks += _text_blocks(body[pos:])
    return blocks


def _braced(s: str, cmd: str) -> str:
    i = s.index(cmd + "{") + len(cmd) + 1
    depth, j = 1, i
    while depth:
        depth += {"{": 1, "}": -1}.get(s[j], 0)
        j += 1
    return s[i : j - 1]


def _text_blocks(chunk: str) -> list[dict]:
    out = []
    for para in re.split(r"\n\s*\n", chunk):
        para = para.strip()
        if not para:
            continue
        m = re.match(r"\\(section|subsection)\{([^}]*)\}\s*(.*)", para, re.S)
        if m:
            out.append({"kind": m.group(1), "title": m.group(2)})
            para = m.group(3).strip()
            if not para:
                continue
        m = re.match(r"\\paragraph\{([^}]*)\}\s*(.*)", para, re.S)
        if m:
            out.append({"kind": "subsection", "title": m.group(1).rstrip(".")})
            para = m.group(2)
        out.append({"kind": "para", "text": para})
    return out


# ---------------------------------------------------------------------------------------------------------------
# Word writing
# ---------------------------------------------------------------------------------------------------------------

_xslt = etree.XSLT(etree.parse(str(XSL))) if XSL.exists() else None


def omml(latex: str):
    mathml = latex_to_mathml(latex)
    if _xslt is None:
        return None
    return etree.fromstring(str(_xslt(etree.fromstring(mathml.encode()))).encode())


def add_segments(p, segs, size: float, italic: bool = False) -> None:
    for kind, val in segs:
        if kind == "t":
            r = p.add_run(val)
            r.font.size = Pt(size)
            r.italic = italic or None
        else:
            el = omml(val)
            if el is None:
                r = p.add_run(val)
                r.font.size = Pt(size)
            else:
                p._p.append(el)


def para(doc, text="", size=10.0, bold=False, italic=False, align=None, first_line=True, after=0, before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(after), Pt(before)
    if first_line:
        pf.first_line_indent = Twips(284)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold or None, italic or None
    return p


def section_break(doc, sect_template) -> None:
    """End the current section at a new paragraph with the given section properties (continuous)."""
    p = doc.add_paragraph()
    sp = copy.deepcopy(sect_template)
    for ref in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
        sp.remove(ref)
    t = sp.find(qn("w:type"))
    if t is None:
        t = etree.SubElement(sp, qn("w:type"))
        sp.insert(0, t)
    t.set(qn("w:val"), "continuous")
    p._p.get_or_add_pPr().append(sp)
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        r.font.size = Pt(2)


def set_cell_border(cell, **edges) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = etree.SubElement(tcPr, qn("w:tcBorders"))
    for edge, val in edges.items():
        el = etree.SubElement(borders, qn(f"w:{edge}"))
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def build_manuscript(blind: bool) -> Path:
    tex = (PAPER / "main.tex").read_text(encoding="utf-8")
    bib = parse_bib((PAPER / "references.bib").read_text(encoding="utf-8"))
    cites = Citations(bib)
    blocks = parse_body(tex)

    labels = {}
    nfig = ntab = 0
    for b in blocks:
        if b["kind"] == "figure":
            nfig += 1
            labels[b["label"]] = str(nfig)
        elif b["kind"] == "table":
            ntab += 1
            labels[b["label"]] = str(ntab)

    doc = docx.Document(str(TEMPLATE))
    body = doc.element.body
    sect_one = copy.deepcopy(next(body.iter(qn("w:sectPr"))))  # first-page, single column
    sect_two = copy.deepcopy(body.find(qn("w:sectPr")))  # two columns
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    C = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, TITLE_EN, 18, align=C, first_line=False, after=4)
    para(doc, TITLE_ES, 14, align=C, first_line=False, after=8)
    if not blind:
        p = para(doc, AUTHOR, 13, italic=True, align=C, first_line=False)
        sup = p.add_run("a")
        sup.font.superscript, sup.font.size = True, Pt(13)
        para(doc, f"a {AFFILIATION}. {EMAIL}. ORCID: {ORCID}", 8, italic=True, align=C, first_line=False, after=8)
    para(doc, "Abstract", 9, bold=True, first_line=False, before=6)
    para(doc, ABSTRACT_EN, 9, first_line=False, after=2)
    p = para(doc, "Keywords: ", 9, italic=True, first_line=False, after=6)
    p.add_run("; ".join(KEYWORDS_EN) + ".").italic = True
    para(doc, "Resumen", 9, bold=True, first_line=False)
    para(doc, ABSTRACT_ES, 9, first_line=False, after=2)
    p = para(doc, "Palabras clave: ", 9, italic=True, first_line=False, after=6)
    p.add_run("; ".join(KEYWORDS_ES) + ".").italic = True
    section_break(doc, sect_one)

    col_width = 5049
    nsec = nsub = neq = 0
    for b in blocks:
        k = b["kind"]
        if k == "section":
            nsec, nsub = nsec + 1, 0
            para(
                doc,
                f"{nsec} {b['title']}",
                10,
                bold=True,
                first_line=False,
                before=8,
                after=4,
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
        elif k == "subsection":
            nsub += 1
            p = para(
                doc,
                f"{nsec}.{nsub} {b['title']}",
                10,
                bold=True,
                italic=True,
                first_line=False,
                before=4,
                after=2,
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
        elif k == "para":
            p = para(doc, first_line=True)
            add_segments(p, latex_text(b["text"], cites, labels), 10)
        elif k == "equation":
            neq += 1
            p = para(doc, first_line=False, before=3, after=3)
            ts = p.paragraph_format.tab_stops
            ts.add_tab_stop(Twips(col_width // 2), WD_TAB_ALIGNMENT.CENTER)
            ts.add_tab_stop(Twips(col_width), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t")
            add_segments(p, [("m", b["latex"])], 10)
            p.add_run(f"\t({neq})").font.size = Pt(10)
        elif k == "figure":
            if b["wide"]:
                section_break(doc, sect_two)
            p = para(doc, first_line=False, align=C, before=6)
            width = Twips(10325) if b["wide"] else Twips(col_width)
            p.add_run().add_picture(str(PAPER / "figures" / b["file"]), width=width)
            cap = para(doc, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            r = cap.add_run(f"Figure {labels[b['label']]}. ")
            r.font.size = Pt(8)
            add_segments(cap, latex_text(b["caption"], cites, labels), 8)
            para(doc, "Source: The author.", 8, first_line=False, after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
            if b["wide"]:
                section_break(doc, sect_one)
        elif k == "table":
            section_break(doc, sect_two)  # tables span both columns
            para(doc, f"Table {labels[b['label']]}.", 8, first_line=False, before=6, align=WD_ALIGN_PARAGRAPH.LEFT)
            cap = para(doc, first_line=False, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
            add_segments(cap, latex_text(b["caption"], cites, labels), 8)
            rows = b["rows"]
            ncol = max(len(r["cells"]) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            table_width = 10325
            first = int(table_width * 0.36)
            widths = [first] + [(table_width - first) // (ncol - 1)] * (ncol - 1)
            for i, row in enumerate(rows):
                for j in range(ncol):
                    cell = t.cell(i, j)
                    cell.width = Twips(widths[j])
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.first_line_indent = 0
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                    text = row["cells"][j] if j < len(row["cells"]) else ""
                    if text:
                        add_segments(cp, latex_text(text, cites, labels), 8)
                    edges = {}
                    if row["rule_before"] in ("\\toprule", "\\midrule"):
                        edges["top"] = "single"
                    if i == len(rows) - 1:
                        edges["bottom"] = "single"
                    if edges:
                        set_cell_border(cell, **edges)
            para(doc, "Source: The author.", 8, first_line=False, after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
            section_break(doc, sect_one)

    repo = BLIND_REPO if blind else REPO
    back = [
        (
            "Data and code availability",
            f"The software (gravanom, API and web interface), the synthetic benchmark, the validation protocol and the "
            f"scripts that generate the figures are openly available at {repo} under the BSD-3-Clause licence "
            f"(documentation and synthetic data: CC BY 4.0). All data used in this article are synthetic and can be "
            f"regenerated from the published seeds.",
        ),
        ("Declaration of generative AI use", AI_STATEMENT),
        (
            "Funding and competing interests",
            "This research did not receive any specific grant from funding agencies in the public, commercial or "
            "not-for-profit sectors. The author declares no competing interests.",
        ),
    ]
    for head, text in back:
        para(doc, head, 10, bold=True, first_line=False, before=8, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
        para(doc, text, 10)

    para(doc, "References", 10, bold=True, first_line=False, before=8, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, key in enumerate(cites.order, 1):
        p = para(doc, f"[{i}] {format_reference(bib[key])}", 8, first_line=False, after=3)
        p.paragraph_format.left_indent = Twips(357)
        p.paragraph_format.first_line_indent = Twips(-357)

    if not blind:
        p = para(doc, "", first_line=False, before=10)
        r = p.add_run("A.J. Herrera, ")
        r.bold, r.font.size = True, Pt(8)
        p.add_run(BIO.removeprefix("A.J. Herrera ")).font.size = Pt(8)
        para(doc, f"ORCID: {ORCID}", 8, first_line=False)

    name = "manuscript_blind.docx" if blind else "manuscript_full.docx"
    doc.save(str(OUT / name))
    unused = sorted(set(bib) - set(cites.order))
    print(
        f"{name}: {nfig} figures, {ntab} tables, {neq} equations, {len(cites.order)} references"
        f" (not cited: {', '.join(unused)})"
    )
    return OUT / name


def build_cover_letter() -> Path:
    doc = docx.Document(str(LETTER_TEMPLATE))
    selected = {"Geoquímica y geofísica", "Aplicaciones de software e inteligencia artificial en ciencias de la Tierra"}
    reviewers = [
        (
            "Carlos Alberto Vargas Jiménez",
            "Departamento de Geociencias, Universidad Nacional de Colombia, Bogotá (nacional)",
        ),
        ("Luis Hernán Ochoa Gutiérrez", "Departamento de Geociencias, Universidad Nacional de Colombia, Bogotá (nacional)"),
        (
            "Leonardo Uieda",
            "Instituto de Astronomia, Geofísica e Ciências Atmosféricas, Universidade de São Paulo, Brasil (internacional)",
        ),
        (
            "Santiago R. Soler",
            "Department of Earth, Ocean and Atmospheric Sciences, University of British Columbia, Canadá (internacional)",
        ),
    ]
    paras = doc.paragraphs
    rev_i = 0
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t == "Fecha y Ciudad":
            _set(p, "Armenia (Quindío), 22 de octubre de 2026")
        elif t.startswith("Los autores firmantes remitimos"):
            _set(
                p,
                f"Remito el manuscrito titulado “{TITLE_EN}” ({TITLE_ES}) para que sea considerado para su "
                f"publicación en la Revista BOLETÍN DE CIENCIAS DE LA TIERRA. Se trata de un artículo de "
                f"investigación de autoría única.",
            )
        elif t.startswith("Que, en caso de haber utilizado herramientas de inteligencia artificial"):
            _set(
                p,
                t + " Declaración: durante la preparación del trabajo utilicé Claude (Anthropic), un asistente de "
                "inteligencia artificial, como apoyo en el desarrollo y las pruebas del software y en la redacción "
                "y edición del texto; revisé todo el contenido y asumo plena responsabilidad por él.",
            )
        elif t.startswith("Nombre: ____") and rev_i < len(reviewers):
            _set(p, f"Nombre: {reviewers[rev_i][0]}")
            _set(paras[i + 1], f"Afiliación: {reviewers[rev_i][1]}")
            rev_i += 1

    # Thematic areas: mark the box to the right of the selected areas.
    for row in doc.tables[0].rows:
        cells = row.cells
        for j in range(0, len(cells) - 1, 2):
            if cells[j].text.strip() in selected:
                _set(cells[j + 1].paragraphs[0], "X")
    # Contribution table.
    contrib = doc.tables[1]
    _set(contrib.cell(0, 1).paragraphs[0], f"{AUTHOR} (autoría única)")
    _set(
        contrib.cell(1, 1).paragraphs[0],
        "Autoría única: concepción del estudio, diseño de la metodología, "
        "desarrollo del software, análisis y validación, figuras y redacción.",
    )
    _set(
        contrib.cell(2, 1).paragraphs[0],
        "Desarrollo del flujo reproducible de reducción e interpolación, "
        "campaña sintética de verificación, validación cruzada "
        "espacial e interpretación de resultados.",
    )
    _set(contrib.cell(3, 1).paragraphs[0], "No aplica: el manuscrito tiene una sola persona autora.")

    # Signature block: fill the first one, drop the two blocks meant for co-authors.
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip().startswith("(Firmado por todos los autores)"))
    sig = [
        p
        for p in paras[start:]
        if p.text.strip().startswith(
            ("Firma:", "Nombre(s) y apellido(s)", "Documento de identidad", "Institución Laboral", "Dirección electrónica")
        )
    ]
    first, rest = sig[:5], sig[5:]
    for p in first:
        t = p.text.strip()
        if t.startswith("Nombre(s)"):
            _set(p, f"Nombre(s) y apellido(s): {AUTHOR}")
        elif t.startswith("Institución Laboral"):
            _set(p, "Institución Laboral: Universidad del Quindío, Armenia, Colombia")
        elif t.startswith("Dirección electrónica"):
            _set(p, f"Dirección electrónica: {EMAIL} — Teléfono: {PHONE} — ORCID: {ORCID}")
    for p in rest:
        p._p.getparent().remove(p._p)
    for p in reversed(doc.paragraphs):
        if p.text.strip() or p._p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            break
        p._p.getparent().remove(p._p)
    path = OUT / "cover_letter.docx"
    doc.save(str(path))
    print("cover_letter.docx written")
    return path


def _set(p, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


if __name__ == "__main__":
    build_manuscript(blind=True)
    build_manuscript(blind=False)
    build_cover_letter()
